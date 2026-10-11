"""Who can see what, and who can do what. These are the most important rules of the app:
- you must be logged in to see anything;
- admins see and do everything, in every project;
- everyone else only sees the projects they manage or whose team they are in.
  Anything in another project answers "Page not found" (404), as if it didn't exist;
- in a project, its manager (and the admins) edit the project, its tasks and its team;
  the members read everything, change the status of their own tasks, comment and add files;
- only admins create projects and choose a project's manager.
The server checks these rules itself: hiding a button is not enough."""

import io
import re
import shutil
import tempfile

from django.contrib.auth.models import AnonymousUser
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import transaction
from django.test import override_settings
from django.urls import URLPattern, URLResolver, get_resolver
from openpyxl import load_workbook

from accounts.models import User
from projects.emails import reminder_for
from projects.models import Attachment, Comment, Project, Task

from .base import TeamTestCase, days_from_today
from .test_activity import visible_text

# What a page answers
OK, MOVED, FORBIDDEN, HIDDEN = 200, 302, 403, 404

# Addresses that aren't about projects: logging in and out, and the admin panel
# (Django lets only admins with "staff status" into the admin panel)
NOT_ABOUT_PROJECTS = {"login", "logout", "admin"}


class LoginRequiredTests(TeamTestCase):
    def test_pages_need_a_login(self):
        pages = [
            "/",
            "/projects/",
            f"/projects/{self.project.pk}/",
            f"/projects/{self.project.pk}/board/",
            f"/projects/{self.project.pk}/timeline/",
            f"/projects/{self.project.pk}/export.xlsx",
            f"/tasks/{self.task.pk}/",
            "/my-tasks/",
            "/activity/",
            "/dashboard/",
            "/reminders/preview/",
        ]
        for page in pages:
            with self.subTest(page=page):  # if one page fails, the message says which one
                self.assertRedirects(self.client.get(page), f"/login/?next={page}")


class EveryAddressTests(TeamTestCase):
    """Every address of the app, opened by five people:
    Amina (admin), Omar (the project's manager), Sara (in the team, the task is hers),
    Lina (in the team, the task isn't hers) and Karim (outside the team).
    The list must name every address of the app: a new page without a line here fails the test."""

    def setUp(self):
        media_folder = tempfile.mkdtemp()  # uploaded files go to a temporary folder
        settings_change = override_settings(MEDIA_ROOT=media_folder)
        settings_change.enable()
        self.addCleanup(settings_change.disable)
        self.addCleanup(shutil.rmtree, media_folder, ignore_errors=True)

        self.lina = User.objects.create_user("lina", "lina@example.com", first_name="Lina")
        self.project.add_to_team(self.lina)
        self.comment = Comment.objects.create(task=self.task, author=self.sara, body="Rack ready")
        self.attachment = Attachment.objects.create(
            task=self.task, uploaded_by=self.sara, file=SimpleUploadedFile("plan.pdf", b"%PDF-1.4 test")
        )

    def rules(self):
        """(name of the address, method, address, data sent,
            what Amina, Omar, Sara, Lina and Karim get)"""
        p, t = self.project.pk, self.task.pk
        dates = {"start_date": days_from_today(1).isoformat(), "due_date": days_from_today(3).isoformat()}
        return [
            ("home", "get", "/", None, (OK, OK, OK, OK, OK)),
            ("my_tasks", "get", "/my-tasks/", None, (OK, OK, OK, OK, OK)),
            ("dashboard", "get", "/dashboard/", None, (OK, OK, OK, OK, OK)),
            ("activity", "get", f"/activity/?project={p}", None, (OK, OK, OK, OK, HIDDEN)),
            ("reminder_preview", "get", "/reminders/preview/", None, (OK, OK, OK, OK, OK)),
            ("project_list", "get", "/projects/", None, (OK, OK, OK, OK, OK)),
            ("projects_excel", "get", "/projects/export.xlsx", None, (OK, OK, OK, OK, OK)),
            ("project_create", "get", "/projects/new/", None, (OK, FORBIDDEN, FORBIDDEN, FORBIDDEN, FORBIDDEN)),
            ("project_detail", "get", f"/projects/{p}/", None, (OK, OK, OK, OK, HIDDEN)),
            ("project_edit", "get", f"/projects/{p}/edit/", None, (OK, OK, FORBIDDEN, FORBIDDEN, HIDDEN)),
            ("team_add", "post", f"/projects/{p}/team/add/", {"person": self.karim.pk},
             (MOVED, MOVED, FORBIDDEN, FORBIDDEN, HIDDEN)),
            ("team_remove", "post", f"/projects/{p}/team/{self.lina.pk}/remove/", None,
             (MOVED, MOVED, FORBIDDEN, FORBIDDEN, HIDDEN)),
            ("project_board", "get", f"/projects/{p}/board/", None, (OK, OK, OK, OK, HIDDEN)),
            ("project_gantt", "get", f"/projects/{p}/timeline/", None, (OK, OK, OK, OK, HIDDEN)),
            ("project_report", "get", f"/projects/{p}/report/", None, (OK, OK, OK, OK, HIDDEN)),
            ("project_tasks_excel", "get", f"/projects/{p}/export.xlsx", None, (OK, OK, OK, OK, HIDDEN)),
            ("task_create", "get", f"/projects/{p}/tasks/new/", None, (OK, OK, FORBIDDEN, FORBIDDEN, HIDDEN)),
            ("task_detail", "get", f"/tasks/{t}/", None, (OK, OK, OK, OK, HIDDEN)),
            ("task_edit", "get", f"/tasks/{t}/edit/", None, (OK, OK, FORBIDDEN, FORBIDDEN, HIDDEN)),
            ("task_status", "post", f"/tasks/{t}/status/", {"status": "in_progress"},
             (MOVED, MOVED, MOVED, FORBIDDEN, HIDDEN)),
            ("task_delete", "get", f"/tasks/{t}/delete/", None, (OK, OK, FORBIDDEN, FORBIDDEN, HIDDEN)),
            ("task_move", "post", f"/tasks/{t}/move/", {"status": "in_progress", "order": [t]},
             (OK, OK, OK, FORBIDDEN, HIDDEN)),
            ("task_dates", "post", f"/tasks/{t}/dates/", dates, (OK, OK, FORBIDDEN, FORBIDDEN, HIDDEN)),
            ("comment_add", "post", f"/tasks/{t}/comments/add/", {"body": "Hello"},
             (MOVED, MOVED, MOVED, MOVED, HIDDEN)),
            ("comment_delete", "post", f"/comments/{self.comment.pk}/delete/", None,  # Sara's comment
             (MOVED, MOVED, MOVED, FORBIDDEN, HIDDEN)),
            ("attachment_add", "post", f"/tasks/{t}/files/add/",
             lambda: {"file": SimpleUploadedFile("quote.pdf", b"%PDF-1.4 quote")},  # a new file each time
             (MOVED, MOVED, MOVED, MOVED, HIDDEN)),
            ("attachment_delete", "post", f"/files/{self.attachment.pk}/delete/", None,  # Sara's file
             (MOVED, MOVED, MOVED, FORBIDDEN, HIDDEN)),
            ("attachment_download", "get", self.attachment.file.url, None, (OK, OK, OK, OK, HIDDEN)),
        ]

    def test_every_address_has_a_rule(self):
        def names(patterns):
            for pattern in patterns:
                if isinstance(pattern, URLResolver):  # a group of addresses, like the admin panel
                    if pattern.app_name == "admin":
                        yield "admin"
                    else:
                        yield from names(pattern.url_patterns)
                elif isinstance(pattern, URLPattern):
                    yield pattern.name

        all_names = set(names(get_resolver().url_patterns))
        checked = {rule[0] for rule in self.rules()}
        self.assertEqual(all_names - NOT_ABOUT_PROJECTS - checked, set(), "These addresses have no rule here")
        self.assertEqual(checked - all_names, set(), "These rules are about addresses that don't exist")

    def test_what_each_person_gets(self):
        people = [self.amina, self.omar, self.sara, self.lina, self.karim]
        for name, method, address, data, expected in self.rules():
            for person, status in zip(people, expected):
                with self.subTest(address=name, person=person.username):
                    self.login(person)
                    # Each request is undone right after, so every person finds the same data
                    with transaction.atomic():
                        send = getattr(self.client, method)
                        response = send(address, data() if callable(data) else data or {})
                        transaction.set_rollback(True)
                    if response.streaming:
                        b"".join(response.streaming_content)  # a downloaded file: read it all, which closes it
                    self.assertEqual(response.status_code, status)


class OnlyYourProjectsTests(TeamTestCase):
    """Lists, numbers and exports only show what is in the person's projects."""

    def setUp(self):
        # Karim manages "Wi-Fi". Sara isn't in its team, so she can't see it.
        self.wifi = Project.objects.create(name="Wi-Fi", manager=self.karim, status="in_progress")
        self.order = Task.objects.create(project=self.wifi, title="Order access points",
                                         assignee=self.karim, due_date=days_from_today(-1))
        Task.objects.create(project=self.wifi, title="Survey", status="done")

    def test_project_list_and_its_export(self):
        for person, expected in [(self.sara, ["Firewall"]), (self.karim, ["Wi-Fi"]),
                                 (self.amina, ["Firewall", "Wi-Fi"])]:
            with self.subTest(person=person.username):
                self.login(person)
                listed = sorted(project.name for project in self.client.get("/projects/").context["projects"])
                self.assertEqual(listed, expected)
                workbook = load_workbook(io.BytesIO(self.client.get("/projects/export.xlsx").content))
                exported = sorted(row[0] for row in workbook["Projects"].iter_rows(min_row=2, values_only=True))
                self.assertEqual(exported, expected)

    def test_home_and_dashboard_numbers(self):
        # What Sara sees (only Firewall), and what Amina sees (Firewall and Wi-Fi)
        for person, projects, overdue, done, attention in [(self.sara, 1, 0, 0, []), (self.amina, 2, 1, 1, ["Wi-Fi"])]:
            with self.subTest(person=person.username):
                self.login(person)
                self.assertEqual(self.client.get("/").context["active_projects"], projects)
                context = self.client.get("/dashboard/").context
                numbers = {kpi["label"]: kpi["value"] for kpi in context["kpis"]}
                self.assertEqual(
                    (numbers["Active projects"], numbers["Overdue tasks"], numbers["Done in the last 30 days"]),
                    (projects, overdue, done),
                )
                self.assertEqual(sum(week["count"] for week in context["done_per_week"]), done)
                self.assertEqual([project.name for project in context["needs_attention"]], attention)
                self.assertEqual(len(context["oldest_overdue"]), overdue)

    def test_tasks_you_cannot_open_are_never_listed(self):
        # By an old mistake, two tasks of Firewall (a project Karim can't see) are his.
        # update() writes straight into the database, without the app's rules.
        Task.objects.filter(pk=self.task.pk).update(assignee=self.karim, due_date=days_from_today(-2))
        done = Task.objects.create(project=self.project, title="Old work", status="done")
        Task.objects.filter(pk=done.pk).update(assignee=self.karim)

        self.login(self.karim)  # only "Order access points", in his own project, may show
        home = self.client.get("/").context
        self.assertEqual((home["open_tasks"], home["overdue_tasks"]), (1, 1))
        my_tasks = self.client.get("/my-tasks/").context
        self.assertEqual(my_tasks["open_count"], 1)
        self.assertEqual(list(my_tasks["recently_done"]), [])
        sections = reminder_for(self.karim, days_from_today(0))["sections"]
        self.assertEqual([task.title for section in sections for task in section.tasks], ["Order access points"])

    def test_activity_feed_and_its_filters(self):
        self.login(self.sara)
        response = self.client.get("/activity/")
        self.assertNotIn("Order access points", visible_text(response))
        self.assertEqual(list(response.context["projects"]), [self.project])
        self.assertEqual(set(response.context["people"]), {self.omar, self.sara, self.amina})  # her team and the admins
        self.assertEqual(self.client.get(f"/activity/?project={self.wifi.pk}").status_code, 404)
        self.assertEqual(self.client.get(f"/activity/?task={self.order.pk}").status_code, 404)
        self.assertEqual(self.client.get(f"/activity/?person={self.karim.pk}").status_code, 404)

        self.login(self.amina)  # admins see everything
        self.assertIn("Order access points", visible_text(self.client.get("/activity/")))

    def test_only_admins_see_what_happened_in_deleted_projects(self):
        self.wifi.delete()
        self.login(self.karim)  # he managed it, but it doesn't exist anymore
        self.assertNotIn("Wi-Fi", visible_text(self.client.get("/activity/")))
        self.login(self.amina)
        self.assertIn("deleted the project Wi-Fi", visible_text(self.client.get("/activity/")))

    def test_leaving_a_team_means_losing_access(self):
        lina = User.objects.create_user("lina", "lina@example.com")
        self.project.add_to_team(lina)
        self.login(lina)
        self.assertEqual(self.client.get(f"/projects/{self.project.pk}/").status_code, 200)
        self.login(self.omar)
        self.client.post(f"/projects/{self.project.pk}/team/{lina.pk}/remove/")
        self.login(lina)
        self.assertEqual(self.client.get(f"/projects/{self.project.pk}/").status_code, 404)
        self.assertEqual(self.client.get(f"/tasks/{self.task.pk}/").status_code, 404)

    def test_nobody_logged_in_sees_nothing(self):
        self.assertFalse(Project.objects.visible_to(AnonymousUser()).exists())
        self.assertFalse(Task.objects.visible_to(AnonymousUser()).exists())

    def test_someone_in_no_team_is_told_why_the_list_is_empty(self):
        lina = User.objects.create_user("lina", "lina@example.com")
        self.login(lina)
        self.assertContains(self.client.get("/projects/"), "You aren't in any project's team yet.")


class NoLeakTests(TeamTestCase):
    """Follow every link of the app as one person, like a visitor clicking everywhere, and check
    that nothing from a project they aren't part of shows anywhere: its name, its tasks, its
    comments, its files. A new page is checked too, as soon as a link leads to it."""

    SECRETS = ["Secret Wi-Fi", "Secret access points", "Secret comment", "secret-quote"]

    def setUp(self):
        media_folder = tempfile.mkdtemp()
        settings_change = override_settings(MEDIA_ROOT=media_folder)
        settings_change.enable()
        self.addCleanup(settings_change.disable)
        self.addCleanup(shutil.rmtree, media_folder, ignore_errors=True)

        # Karim's project, with a team that doesn't include Sara
        wifi = Project.objects.create(name="Secret Wi-Fi", manager=self.karim, status="in_progress")
        task = Task.objects.create(project=wifi, title="Secret access points", assignee=self.karim,
                                   due_date=days_from_today(-1))
        Comment.objects.create(task=task, author=self.karim, body="Secret comment")
        Attachment.objects.create(task=task, uploaded_by=self.karim,
                                  file=SimpleUploadedFile("secret-quote.pdf", b"%PDF-1.4"))
        # Sara's project has some of everything too, so the visit goes through every kind of page
        Comment.objects.create(task=self.task, author=self.sara, body="Rack ready")
        Attachment.objects.create(task=self.task, uploaded_by=self.sara,
                                  file=SimpleUploadedFile("plan.pdf", b"%PDF-1.4"))

    def visit_everything(self, person):
        """Every page reachable from the home page by following links: {address: text}."""
        self.login(person)
        # The home page, and two lists that are opened by a form (not by a link)
        pages, to_visit = {}, ["/", "/projects/?status=all", "/projects/export.xlsx?status=all"]
        while to_visit and len(pages) < 400:
            address = to_visit.pop()
            if address in pages:
                continue
            response = self.client.get(address)
            if response.streaming:  # a downloaded file
                pages[address] = b"".join(response.streaming_content).decode("utf-8", "replace")
                continue
            if response.status_code != 200:
                pages[address] = ""
                continue
            if "spreadsheetml" in response["Content-Type"]:  # an Excel file: read its cells
                workbook = load_workbook(io.BytesIO(response.content))
                pages[address] = " ".join(str(cell) for sheet in workbook
                                          for row in sheet.iter_rows(values_only=True) for cell in row)
                continue
            pages[address] = text = response.content.decode()
            for link in re.findall(r'href="(/[^"#]*)', text):
                link = link.replace("&amp;", "&")
                if not link.startswith(("/static/", "/logout/", "/admin/")):
                    to_visit.append(link)
        return pages

    def test_nothing_from_another_project_shows_anywhere(self):
        pages = self.visit_everything(self.sara)
        # The visit really went everywhere: lists, the project's pages, the task, the files...
        p, t = self.project.pk, self.task.pk
        for address in ["/my-tasks/", "/dashboard/", "/activity/", "/reminders/preview/", f"/projects/{p}/board/",
                        f"/projects/{p}/timeline/", f"/projects/{p}/report/", f"/projects/{p}/export.xlsx",
                        f"/tasks/{t}/", f"/activity/?task={t}"]:
            self.assertIn(address, pages)
        self.assertTrue(any(address.startswith("/media/") for address in pages))
        for address, text in pages.items():
            for secret in self.SECRETS:
                self.assertNotIn(secret, text, f"{address} shows {secret!r}")

    def test_the_visit_would_notice(self):
        # The same visit, by someone in Karim's project: then the secrets are found
        pages = self.visit_everything(self.karim)
        for secret in self.SECRETS:
            self.assertTrue(any(secret in text for text in pages.values()), f"{secret!r} was never seen")


class ProjectRightsTests(TeamTestCase):
    def test_only_admins_create_projects(self):
        new_project = {"name": "Wi-Fi", "status": "planned", "priority": "high", "manager": self.sara.pk}
        self.login(self.omar)  # he manages a project, but isn't an admin
        self.assertEqual(self.client.post("/projects/new/", new_project).status_code, 403)
        self.assertFalse(Project.objects.filter(name="Wi-Fi").exists())

        self.login(self.amina)
        response = self.client.post("/projects/new/", new_project)
        wifi = Project.objects.get(name="Wi-Fi")
        self.assertRedirects(response, f"/projects/{wifi.pk}/")
        self.assertEqual(wifi.manager, self.sara)  # any user can be a project's manager
        self.login(self.sara)
        self.assertEqual(self.client.get(f"/projects/{wifi.pk}/edit/").status_code, 200)

    def test_the_manager_edits_the_project_but_not_who_manages_it(self):
        self.login(self.omar)
        response = self.client.get(f"/projects/{self.project.pk}/edit/")
        self.assertContains(response, "Only admins can change the manager.")
        self.assertEqual(list(response.context["form"].fields["manager"].queryset), [self.omar])  # nothing else
        self.client.post(f"/projects/{self.project.pk}/edit/", {
            "name": "Firewall 2", "status": "in_progress", "priority": "high", "manager": self.karim.pk})
        self.project.refresh_from_db()
        self.assertEqual((self.project.name, self.project.manager), ("Firewall 2", self.omar))

    def test_members_cannot_edit_the_project(self):
        self.login(self.sara)
        data = {"name": "Mine now", "status": "in_progress", "priority": "low", "manager": self.sara.pk}
        self.assertEqual(self.client.post(f"/projects/{self.project.pk}/edit/", data).status_code, 403)
        self.project.refresh_from_db()
        self.assertEqual(self.project.name, "Firewall")


class TaskRightsTests(TeamTestCase):
    def test_members_cannot_create_edit_or_delete_tasks(self):
        self.login(self.sara)
        self.assertEqual(self.client.get(f"/projects/{self.project.pk}/tasks/new/").status_code, 403)
        self.assertEqual(self.client.post(f"/tasks/{self.task.pk}/edit/", self.task_form_data(title="Mine now")).status_code, 403)
        self.assertEqual(self.client.post(f"/tasks/{self.task.pk}/delete/").status_code, 403)
        self.task.refresh_from_db()
        self.assertEqual(self.task.title, "Install the firewall")  # nothing changed

    def test_the_project_manager_can_edit_and_delete_tasks(self):
        self.login(self.omar)
        self.client.post(f"/tasks/{self.task.pk}/edit/", self.task_form_data(title="Install it"))
        self.task.refresh_from_db()
        self.assertEqual(self.task.title, "Install it")

        response = self.client.post(f"/tasks/{self.task.pk}/delete/")
        self.assertRedirects(response, f"/projects/{self.project.pk}/")
        self.assertFalse(Task.objects.filter(pk=self.task.pk).exists())

    def test_the_assignee_can_change_the_status(self):
        self.login(self.sara)
        self.client.post(f"/tasks/{self.task.pk}/status/", {"status": "in_progress"})
        self.task.refresh_from_db()
        self.assertEqual(self.task.status, "in_progress")

    def test_another_member_cannot_change_the_status(self):
        self.project.add_to_team(self.karim)  # in the team, but the task is Sara's
        self.login(self.karim)
        self.assertEqual(self.client.post(f"/tasks/{self.task.pk}/status/", {"status": "done"}).status_code, 403)
        response = self.client.post(f"/tasks/{self.task.pk}/move/", {"status": "done", "order": [self.task.pk]})
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.json()["error"], "You can only move tasks assigned to you.")
        self.task.refresh_from_db()
        self.assertEqual(self.task.status, "todo")

    def test_only_the_project_manager_changes_dates_on_the_timeline(self):
        url = f"/tasks/{self.task.pk}/dates/"
        dates = {"start_date": days_from_today(1).isoformat(), "due_date": days_from_today(3).isoformat()}
        self.login(self.sara)
        self.assertEqual(self.client.post(url, dates).status_code, 403)
        self.login(self.omar)
        self.assertEqual(self.client.post(url, dates).status_code, 200)
        self.task.refresh_from_db()
        self.assertEqual(self.task.due_date, days_from_today(3))

    def test_only_admins_see_someone_elses_reminder(self):
        self.login(self.omar)
        self.assertContains(self.client.get(f"/reminders/preview/?user={self.sara.pk}"), "Hello Omar")
        self.login(self.amina)
        self.assertContains(self.client.get(f"/reminders/preview/?user={self.sara.pk}"), "Hello Sara")


class CommentPermissionTests(TeamTestCase):
    def test_the_team_can_comment(self):
        self.login(self.sara)
        self.client.post(f"/tasks/{self.task.pk}/comments/add/", {"body": "I can help."})
        self.assertTrue(Comment.objects.filter(task=self.task, author=self.sara, body="I can help.").exists())

    def test_people_delete_their_own_comments_only(self):
        self.project.add_to_team(self.karim)
        comment = Comment.objects.create(task=self.task, author=self.sara, body="Done soon")
        self.login(self.karim)
        self.assertEqual(self.client.post(f"/comments/{comment.pk}/delete/").status_code, 403)
        self.login(self.sara)
        self.client.post(f"/comments/{comment.pk}/delete/")
        self.assertFalse(Comment.objects.filter(pk=comment.pk).exists())

    def test_the_project_manager_can_delete_any_comment(self):
        comment = Comment.objects.create(task=self.task, author=self.sara, body="Off topic")
        self.login(self.omar)
        self.client.post(f"/comments/{comment.pk}/delete/")
        self.assertFalse(Comment.objects.filter(pk=comment.pk).exists())
