"""Languages (Stage 17c): English (the default), French and Arabic (written from right to left).

- which language each page is shown in, and the language switch (and its safety);
- the language of the emails and of the Excel files;
- every page opens in French and in Arabic;
- the translation files: complete, in step with each other, and compiled.
The texts people read in French and Arabic are in locale/ (see scripts/compile_translations.py)."""

import gettext as gettext_module
import io
import os
import re
import subprocess
import sys
import tempfile
from datetime import date, timedelta
from pathlib import Path

from django.conf import settings
from django.contrib.auth.models import AnonymousUser
from django.contrib.staticfiles import finders
from django.core import mail
from django.core.management import call_command
from django.http import HttpResponse
from django.test import Client, RequestFactory, SimpleTestCase, TestCase, override_settings
from django.utils import timezone, translation
from django.utils.formats import date_format
from django.utils.timesince import TIME_STRINGS, timesince
from django.utils.translation import gettext, ngettext
from openpyxl import load_workbook

from accounts.middleware import LanguageMiddleware, chosen_language
from accounts.models import User
from projects import activity
from projects.activity import Event
from projects.emails import email_language, reminder_for
from projects.models import Comment, Task
from scripts.compile_translations import TranslationError, compile_po, read_po

from .base import TeamTestCase, days_from_today

COOKIE = settings.LANGUAGE_COOKIE_NAME  # the cookie where this computer remembers the language
LOCALE_FOLDER = Path(settings.BASE_DIR) / "locale"


def choose(person, language):
    """Saves a language in someone's profile, as the language switch does."""
    User.objects.filter(pk=person.pk).update(language=language)


class WhichLanguageTests(TeamTestCase):
    """The language of a page: the one in the person's profile, else the one chosen on this computer
    (for the login page), else English. The language of the browser is never used."""

    def test_english_is_the_default(self):
        self.login(self.sara)  # Sara has not chosen a language
        response = self.client.get("/")
        self.assertContains(response, '<html lang="en" dir="ltr">')
        self.assertEqual(response["Content-Language"], "en")
        self.assertContains(response, "Log out")
        self.assertContains(response, "vendor/tabler/tabler.min.css")
        self.assertNotContains(response, "tabler.rtl")

    def test_the_language_of_the_browser_is_not_used(self):
        response = self.client.get("/login/", HTTP_ACCEPT_LANGUAGE="fr-FR,fr;q=0.9,ar;q=0.8")
        self.assertContains(response, '<html lang="en"')
        self.login(self.sara)
        self.assertContains(self.client.get("/", HTTP_ACCEPT_LANGUAGE="ar"), '<html lang="en"')

    def test_each_person_sees_the_language_of_their_profile(self):
        choose(self.sara, "fr")
        choose(self.omar, "ar")
        self.login(self.sara)
        french = self.client.get("/")
        self.assertContains(french, '<html lang="fr" dir="ltr">')
        self.assertContains(french, "Se déconnecter")
        self.assertEqual(french["Content-Language"], "fr")

        # The browser now remembers "fr" for the login page, but Omar's own profile wins
        self.login(self.omar)
        arabic = self.client.get("/")
        self.assertContains(arabic, '<html lang="ar" dir="rtl">')
        self.assertContains(arabic, "تسجيل الخروج")
        self.assertEqual(arabic["Content-Language"], "ar")

    def test_arabic_is_written_from_right_to_left(self):
        choose(self.sara, "ar")
        self.login(self.sara)
        response = self.client.get("/")
        self.assertContains(response, 'dir="rtl"')
        self.assertContains(response, "vendor/tabler/tabler.rtl.min.css")  # Tabler's mirrored design
        self.assertNotContains(response, "vendor/tabler/tabler.min.css")
        # The timeline stays left to right, as time runs that way in every language
        timeline = self.client.get(f"/projects/{self.project.pk}/timeline/")
        self.assertContains(timeline, '<div id="gantt" dir="ltr"')
        # The mirrored design file really is in the project
        self.assertIsNotNone(finders.find("vendor/tabler/tabler.rtl.min.css"))

    def test_a_language_that_is_not_in_the_list_means_english(self):
        choose(self.sara, "de")
        self.login(self.sara)
        self.assertContains(self.client.get("/"), '<html lang="en"')

    def test_the_rules_in_short(self):
        request = RequestFactory().get("/")
        request.user = AnonymousUser()
        self.assertEqual(chosen_language(request), "en")  # nothing chosen
        request.COOKIES[COOKIE] = "fr"
        self.assertEqual(chosen_language(request), "fr")  # chosen on this computer
        request.COOKIES[COOKIE] = "klingon"
        self.assertEqual(chosen_language(request), "en")  # not one of ours
        request.COOKIES[COOKIE] = "fr"
        request.user = User(language="ar")
        self.assertEqual(chosen_language(request), "ar")  # the profile comes first
        request.user = User(language="")
        self.assertEqual(chosen_language(request), "fr")  # no choice in the profile yet

    def test_the_language_is_put_back_after_each_page(self):
        # A page in Arabic must not leave the next page (or an email) in Arabic
        choose(self.sara, "ar")
        self.login(self.sara)
        before = translation.get_language()
        self.client.get("/")
        self.assertEqual(translation.get_language(), before)

    def test_the_browser_remembers_the_language_of_the_profile(self):
        choose(self.sara, "ar")
        self.login(self.sara)
        response = self.client.get("/")
        self.assertEqual(response.cookies[COOKIE].value, "ar")  # so the login page of this computer follows
        self.assertNotIn(COOKIE, self.client.get("/").cookies)  # up to date: not sent again

    def test_pages_are_marked_as_depending_on_the_cookies(self):
        # The page changes with who is logged in and with the language cookie: a shared cache (a proxy)
        # must not hand one person's page, in one language, to someone else.
        def varies_on(response):
            return [name.strip() for name in response["Vary"].split(",")]

        self.assertIn("Cookie", varies_on(self.client.get("/login/")))
        self.login(self.sara)
        self.assertIn("Cookie", varies_on(self.client.get("/")))
        self.assertIn("Cookie", varies_on(self.client.get("/there-is-nothing-here/")))
        self.assertIn("Cookie", varies_on(self.client.post("/language/", {"language": "fr", "next": "/"})))

    def test_the_middleware_itself_marks_the_page_as_depending_on_the_cookies(self):
        # Without the rest of the app: a page that never reads the session still varies with the language cookie
        request = RequestFactory().get("/")
        request.user = AnonymousUser()
        response = LanguageMiddleware(lambda _request: HttpResponse("hello"))(request)
        self.assertEqual(response["Vary"], "Cookie")
        self.assertEqual(response["Content-Language"], "en")

    def test_error_pages_are_in_the_language_too(self):
        choose(self.sara, "fr")
        self.login(self.sara)
        self.assertContains(self.client.get("/there-is-nothing-here/"), "Page introuvable", status_code=404)
        # Sara is a member of the project, not its manager
        self.assertContains(self.client.get(f"/projects/{self.project.pk}/edit/"), "Accès refusé", status_code=403)
        choose(self.sara, "ar")
        self.assertContains(self.client.get("/there-is-nothing-here/"), "الصفحة غير موجودة", status_code=404)


class LanguageSwitchTests(TeamTestCase):
    """The language buttons (in the user menu, and on the login page)."""

    def switch(self, language, next_page="/projects/"):
        return self.client.post("/language/", {"language": language, "next": next_page})

    def test_the_choice_is_kept_in_the_profile_and_in_the_browser(self):
        self.login(self.sara)
        response = self.switch("fr")
        self.assertRedirects(response, "/projects/")  # the same page opens again
        self.sara.refresh_from_db()
        self.assertEqual(self.sara.language, "fr")
        cookie = response.cookies[COOKIE]
        self.assertEqual((cookie.value, cookie["httponly"], cookie["samesite"], cookie["max-age"]),
                         ("fr", True, "Lax", settings.LANGUAGE_COOKIE_AGE))
        for page in ("/projects/", "/my-tasks/", "/"):  # every page, not only this one
            self.assertContains(self.client.get(page), '<html lang="fr"')

    def test_the_choice_changes_nobody_else(self):
        choose(self.omar, "ar")
        self.login(self.sara)
        self.switch("fr")
        self.assertEqual(User.objects.get(pk=self.omar.pk).language, "ar")  # kept
        self.assertEqual(User.objects.get(pk=self.karim.pk).language, "")  # still no choice
        self.client.logout()
        self.switch("ar")  # nobody is logged in: only the browser remembers
        self.assertEqual(sorted(User.objects.values_list("username", "language").exclude(language="")),
                         [("omar", "ar"), ("sara", "fr")])

    def test_changing_back_to_english(self):
        choose(self.sara, "fr")
        self.login(self.sara)
        self.switch("en")
        self.sara.refresh_from_db()
        self.assertEqual(self.sara.language, "en")  # a real choice, not "nothing chosen yet"
        self.assertContains(self.client.get("/"), '<html lang="en"')

    def test_before_logging_in_the_choice_is_only_kept_in_the_browser(self):
        response = self.switch("ar", "/login/")
        self.assertRedirects(response, "/login/")
        self.assertEqual(response.cookies[COOKIE].value, "ar")
        self.assertFalse(User.objects.exclude(language="").exists())  # nobody's profile changed
        login_page = self.client.get("/login/")
        self.assertContains(login_page, '<html lang="ar" dir="rtl">')
        self.assertContains(login_page, "تسجيل الدخول")

    def test_a_language_that_does_not_exist_changes_nothing(self):
        self.login(self.sara)
        self.client.get("/")  # the browser now has its cookie ("en")
        for wrong in ("de", "", "FR", "fr ", "../etc", "ar;x"):
            with self.subTest(language=wrong):
                response = self.switch(wrong)
                self.assertEqual(response.status_code, 302)
                self.assertNotIn(COOKIE, response.cookies)
        response = self.client.post("/language/", {"next": "/projects/"})  # no language at all
        self.assertNotIn(COOKIE, response.cookies)
        self.sara.refresh_from_db()
        self.assertEqual(self.sara.language, "en")  # as saved when she logged in (nothing chosen yet)

    def test_the_person_only_goes_back_to_this_site(self):
        self.login(self.sara)
        for elsewhere in ("https://evil.example/", "//evil.example/path", "javascript:alert(1)",
                          "/\\evil.example", "http://evil.example", "https:evil.example", ""):
            with self.subTest(next=elsewhere):
                self.assertRedirects(self.switch("fr", elsewhere), "/", fetch_redirect_response=False)
        # Addresses of this site are fine
        self.assertRedirects(self.switch("fr", "/projects/?status=in_progress"), "/projects/?status=in_progress",
                             fetch_redirect_response=False)
        self.assertEqual(self.switch("fr", "http://testserver/projects/").url, "http://testserver/projects/")
        # On the server (https), the person is not sent back to an unencrypted http:// address
        https = self.client.post("/language/", {"language": "fr", "next": "http://testserver/projects/"}, secure=True)
        self.assertRedirects(https, "/", fetch_redirect_response=False)
        https = self.client.post("/language/", {"language": "fr", "next": "https://testserver/projects/"}, secure=True)
        self.assertEqual(https.url, "https://testserver/projects/")
        # A new line in the address can't add a second header to the answer
        response = self.switch("fr", "/projects/\r\nSet-Cookie: stolen=1")
        self.assertNotIn("\n", response["Location"])
        self.assertNotIn("stolen", response.cookies)

    def test_the_switch_needs_a_post(self):
        self.login(self.sara)
        self.assertEqual(self.client.get("/language/?language=fr&next=/projects/").status_code, 405)
        self.sara.refresh_from_db()
        self.assertEqual(self.sara.language, "en")  # unchanged

    def test_the_switch_needs_the_security_token(self):
        # Another website must not be able to change someone's language with a hidden form
        strict = Client(enforce_csrf_checks=True)
        strict.force_login(self.sara)
        self.assertEqual(strict.post("/language/", {"language": "fr", "next": "/"}).status_code, 403)
        self.sara.refresh_from_db()
        self.assertEqual(self.sara.language, "en")  # unchanged

    @override_settings(LANGUAGE_COOKIE_SECURE=True)
    def test_on_the_server_the_cookie_only_travels_encrypted(self):
        self.assertTrue(self.switch("fr", "/login/").cookies[COOKIE]["secure"])

    def test_on_the_server_the_settings_keep_every_cookie_encrypted(self):
        # The settings read the environment once, when the app starts: so look at them in a fresh Python
        code = (
            "import sys; sys.path.insert(0, sys.argv[1]); import django; django.setup(); "
            "from django.conf import settings as s; "
            "print(s.LANGUAGE_COOKIE_SECURE, s.SESSION_COOKIE_SECURE, s.CSRF_COOKIE_SECURE, s.LANGUAGE_COOKIE_HTTPONLY)"
        )
        environment = {**os.environ, "HTTPS": "True", "DJANGO_SETTINGS_MODULE": "config.settings"}
        result = subprocess.run([sys.executable, "-c", code, str(settings.BASE_DIR)], env=environment,
                                capture_output=True, text=True, timeout=60)
        self.assertEqual(result.stdout.strip(), "True True True True", result.stderr[-500:])

    def test_every_page_has_the_three_buttons(self):
        self.login(self.sara)
        for page in ("/", "/projects/", f"/projects/{self.project.pk}/", "/profile/"):
            response = self.client.get(page)
            for name in ("English", "Français", "العربية"):
                self.assertContains(response, name)
            self.assertContains(response, 'action="/language/"')
        self.client.logout()
        login_page = self.client.get("/login/")  # also before logging in
        for name in ("English", "Français", "العربية"):
            self.assertContains(login_page, name)


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])  # fast, for tests only
class LanguageAtLoginTests(TestCase):
    """The first time someone logs in, the language of the login page becomes their language."""

    @classmethod
    def setUpTestData(cls):
        cls.nour = User.objects.create_user("nour", "nour@example.com", "a-Strong-pass-42", first_name="Nour")

    def log_in(self, password="a-Strong-pass-42"):
        return self.client.post("/login/", {"username": "nour", "password": password})

    def test_the_first_login_keeps_the_language_of_the_login_page(self):
        self.client.post("/language/", {"language": "fr", "next": "/login/"})  # chosen on the login page
        self.assertRedirects(self.log_in(), "/", fetch_redirect_response=False)
        self.nour.refresh_from_db()
        self.assertEqual(self.nour.language, "fr")
        self.assertContains(self.client.get("/"), '<html lang="fr"')

    def test_without_a_choice_the_first_login_saves_english(self):
        self.log_in()
        self.nour.refresh_from_db()
        self.assertEqual(self.nour.language, "en")

    def test_a_language_already_in_the_profile_wins_over_the_login_page(self):
        choose(self.nour, "ar")
        self.client.post("/language/", {"language": "fr", "next": "/login/"})
        self.log_in()
        self.nour.refresh_from_db()
        self.assertEqual(self.nour.language, "ar")
        self.assertContains(self.client.get("/"), '<html lang="ar"')
        self.assertEqual(self.client.cookies[COOKIE].value, "ar")  # the login page of this computer follows

    def test_a_wrong_password_saves_nothing(self):
        self.client.post("/language/", {"language": "fr", "next": "/login/"})
        with self.assertLogs("axes", level="WARNING"):  # the failed login is written in the log
            self.log_in("wrong")
        self.nour.refresh_from_db()
        self.assertEqual(self.nour.language, "")


class ProfileLanguageTests(TeamTestCase):
    """My profile page also has the language (Stage 17c)."""

    def save_profile(self, **data):
        return self.client.post("/profile/", {"phone": "", **data}, follow=True)

    def test_my_profile_offers_the_three_languages(self):
        self.login(self.sara)
        page = self.client.get("/profile/")
        for code in ("en", "fr", "ar"):
            self.assertContains(page, f'<option value="{code}"')
        self.assertContains(page, '<option value="en" selected>')  # none chosen yet: English is shown

    def test_choosing_another_language(self):
        self.login(self.sara)
        response = self.save_profile(language="ar")
        self.sara.refresh_from_db()
        self.assertEqual(self.sara.language, "ar")
        self.assertContains(response, '<html lang="ar" dir="rtl">')
        self.assertContains(response, "تم حفظ ملفك الشخصي.")  # the message is already in the new language
        self.assertEqual(self.client.cookies[COOKIE].value, "ar")

    def test_saving_something_else_keeps_the_language(self):
        choose(self.sara, "fr")
        self.login(self.sara)
        self.save_profile(phone="22 12 34 56")  # a form sent without the language
        self.sara.refresh_from_db()
        self.assertEqual((self.sara.phone, self.sara.language), ("22 12 34 56", "fr"))

    def test_an_unknown_language_is_refused(self):
        self.login(self.sara)
        self.assertContains(self.save_profile(language="de"), "Select a valid choice")
        self.sara.refresh_from_db()
        self.assertEqual(self.sara.language, "en")  # unchanged

    def test_only_the_person_changes_their_language(self):
        # The language is not in the Administration forms: nobody picks the language of someone else
        self.login(self.amina)
        self.client.post(f"/administration/users/{self.sara.pk}/", {
            "username": "sara", "first_name": "Sara", "last_name": "Ben", "email": "sara@example.com",
            "phone": "22 12 34 56", "role": "user", "language": "ar",
        })
        self.sara.refresh_from_db()
        self.assertEqual(self.sara.phone, "22 12 34 56")  # the form was accepted...
        self.assertEqual(self.sara.language, "")  # ...but it can't choose her language


@override_settings(SITE_URL="https://itpm.example.org")
@override_settings(EMAIL_DEFAULT_LANGUAGE="en")  # these tests are about English wording (a blank profile = English)
class EmailLanguageTests(TeamTestCase):
    """Each email is written in the language of the person who gets it, whoever caused it."""

    def setUp(self):
        Task.objects.create(project=self.project, title="Fix the VPN", assignee=self.sara,
                            due_date=days_from_today(-2))
        Task.objects.create(project=self.project, title="Renew the licence", assignee=self.sara,
                            due_date=days_from_today(0))
        choose(self.sara, "fr")
        choose(self.omar, "ar")  # Omar manages the project: he is told about the late task

    def test_the_reminder_is_written_in_the_language_of_the_person(self):
        before = translation.get_language()
        french = reminder_for(User.objects.get(pk=self.sara.pk), days_from_today(0))
        self.assertEqual([section.title for section in french["sections"]],
                         ["En retard", "Échéance dans les 3 prochains jours"])
        self.assertEqual(french["subject"], "Vos tâches\xa0: 1 en retard, 1 bientôt à échéance")
        self.assertEqual([task.due_text for section in french["sections"] for task in section.tasks],
                         ["2 jours de retard", "échéance aujourd'hui"])

        arabic = reminder_for(User.objects.get(pk=self.omar.pk), days_from_today(0))
        self.assertEqual([section.title for section in arabic["sections"]], ["متأخرة في المشاريع التي تديرها"])
        self.assertEqual(arabic["subject"], "مهامك: 1 متأخرة في مشاريعك")
        self.assertEqual(translation.get_language(), before)  # the language is put back

    def test_the_command_writes_each_email_in_the_language_of_the_person(self):
        call_command("send_reminders", stdout=io.StringIO())
        emails = {email.to[0]: email for email in mail.outbox}
        self.assertEqual(sorted(emails), ["omar@example.com", "sara@example.com"])

        french = emails["sara@example.com"]
        self.assertEqual(french.subject, "Vos tâches : 1 en retard, 1 bientôt à échéance")
        for text in ("Bonjour Sara,", "2 jours de retard", "Ouvrir mes tâches\xa0: https://itpm.example.org/my-tasks/"):
            self.assertIn(text, french.body)  # French puts a (non-breaking) space before a colon
        self.assertIn('<html lang="fr" dir="ltr">', french.alternatives[0][0])

        arabic = emails["omar@example.com"]
        self.assertEqual(arabic.subject, "مهامك: 1 متأخرة في مشاريعك")
        self.assertIn("مرحبًا Omar،", arabic.body)
        html = arabic.alternatives[0][0]
        self.assertIn('<html lang="ar" dir="rtl">', html)
        self.assertIn("text-align:right", html)  # the email reads from right to left
        self.assertIn("Fix the VPN", html)  # the titles are the people's own words: they are not translated

    def test_a_terminal_that_cannot_show_arabic_does_not_stop_the_command(self):
        # A log file or a terminal of a Windows computer may only know simple letters. The command used to stop
        # on the first Arabic subject, and the people after that one got no email.
        raw = io.BytesIO()
        terminal = io.TextIOWrapper(raw, encoding="ascii", line_buffering=True)
        call_command("send_reminders", stdout=terminal)
        terminal.flush()
        self.assertEqual(sorted(email.to[0] for email in mail.outbox), ["omar@example.com", "sara@example.com"])
        self.assertIn(b'omar <omar@example.com>: sent "', raw.getvalue())
        self.assertIn(b"2 emails sent.", raw.getvalue())

    def test_the_new_task_email_follows_the_new_assignee(self):
        self.login(self.omar)
        for language, subject, direction, button in [
            ("fr", "Nouvelle tâche pour vous : Configure the VPN", "ltr", "Ouvrir la tâche"),
            ("ar", "مهمة جديدة لك: Configure the VPN", "rtl", "فتح المهمة"),
            ("en", "New task for you: Configure the VPN", "ltr", "Open the task"),
        ]:
            with self.subTest(language=language):
                choose(self.sara, language)
                mail.outbox.clear()
                Task.objects.filter(title="Configure the VPN").delete()
                # Omar is logged in with his own language (Arabic): his page doesn't change the email's
                response = self.client.post(f"/projects/{self.project.pk}/tasks/new/", {
                    "title": "Configure the VPN", "status": "todo", "priority": "high",
                    "assignee": self.sara.pk, "due_date": days_from_today(4)}, follow=True)
                self.assertContains(response, '<html lang="ar" dir="rtl">')
                self.assertEqual(len(mail.outbox), 1)
                email = mail.outbox[0]
                self.assertEqual(email.subject, subject)
                self.assertIn(button, email.body)
                self.assertIn(f'<html lang="{language}" dir="{direction}">', email.alternatives[0][0])

    def test_the_preview_shows_the_email_as_its_person_will_get_it(self):
        self.login(self.amina)  # Amina reads English, Sara's reminder is in French
        page = self.client.get(f"/reminders/preview/?user={self.sara.pk}")
        self.assertContains(page, '<html lang="fr"')
        self.assertContains(self.client.get(f"/reminders/preview/?user={self.sara.pk}&text=1"), "Bonjour Sara,")


class ExcelLanguageTests(TeamTestCase):
    """An Excel file is written in the language of the person who downloads it."""

    def download(self, language):
        choose(self.sara, language)
        self.login(self.sara)
        response = self.client.get(f"/projects/{self.project.pk}/export.xlsx")
        self.assertEqual(response.status_code, 200)
        response["Content-Disposition"].encode("ascii")  # the file name stays simple letters, in every language
        return load_workbook(io.BytesIO(response.content))

    def test_english(self):
        book = self.download("en")
        self.assertEqual(book.sheetnames, ["Summary", "Tasks"])
        self.assertFalse(book["Tasks"].sheet_view.rightToLeft)

    def test_french(self):
        Task.objects.create(project=self.project, title="Nobody's")
        book = self.download("fr")
        self.assertEqual(book.sheetnames, ["Résumé", "Tâches"])
        rows = list(book["Tâches"].iter_rows(values_only=True))
        self.assertEqual(rows[0][:5], ("Tâche", "Tâche principale", "Statut", "Priorité", "Assignée à"))
        self.assertIn("Non assignée", [row[4] for row in rows[1:]])
        self.assertFalse(book["Tâches"].sheet_view.rightToLeft)

    def test_arabic_sheets_open_from_right_to_left(self):
        book = self.download("ar")
        self.assertEqual(book.sheetnames, ["ملخص", "المهام"])
        self.assertTrue(all(sheet.sheet_view.rightToLeft for sheet in book))
        self.assertEqual(book["المهام"]["A1"].value, "المهمة")
        self.assertEqual(book["المهام"]["A2"].value, "Install the firewall")  # a title is the people's own word

    def test_the_list_of_projects_too(self):
        choose(self.sara, "fr")
        self.login(self.sara)
        book = load_workbook(io.BytesIO(self.client.get("/projects/export.xlsx").content))
        self.assertEqual(book.sheetnames, ["Projets"])


@override_settings(EMAIL_DEFAULT_LANGUAGE="en")  # these tests are about English wording (a blank profile = English)
class PagesInEveryLanguageTests(TeamTestCase):
    """Every page opens in French and in Arabic, with the right direction and no broken text."""

    def setUp(self):
        # Richer data, so the pages show lists, counts (plural forms), dates and links
        for number in range(1, 12):
            Task.objects.create(project=self.project, title=f"Task {number}", assignee=self.sara,
                                due_date=days_from_today(number - 6))
        self.subtask = Task.objects.create(project=self.project, title="Open the rack", parent=self.task,
                                           assignee=self.sara, due_date=days_from_today(2))
        self.task.depends_on.add(Task.objects.get(title="Task 1"))
        Task.objects.filter(title="Task 2").update(is_milestone=True)
        Comment.objects.create(task=self.task, author=self.omar, body="Rack ready")

    def pages(self):
        p, t, sara = self.project.pk, self.task.pk, self.sara.pk
        return [
            "/", "/projects/", "/projects/?status=in_progress", "/projects/new/",
            f"/projects/{p}/", f"/projects/{p}/edit/", f"/projects/{p}/board/", f"/projects/{p}/timeline/",
            f"/projects/{p}/report/", f"/projects/{p}/delete/", f"/projects/{p}/tasks/new/",
            f"/tasks/{t}/", f"/tasks/{t}/edit/", f"/tasks/{t}/delete/", f"/tasks/{self.subtask.pk}/",
            "/my-tasks/", "/dashboard/", "/activity/", f"/activity/?project={p}",
            "/profile/", "/profile/password/", f"/people/{sara}/", "/reminders/preview/",
            "/administration/", "/administration/users/", "/administration/users/new/",
            f"/administration/users/{sara}/", f"/administration/users/{sara}/password/",
            f"/administration/users/{sara}/deactivate/", "/administration/projects/", "/administration/tasks/",
            "/administration/logins/",
        ]

    def test_every_page_opens_in_french_and_arabic(self):
        self.login(self.amina)  # an admin can open every page
        for language, direction in (("fr", "ltr"), ("ar", "rtl")):
            choose(self.amina, language)
            for page in self.pages():
                with self.subTest(language=language, page=page):
                    response = self.client.get(page)
                    self.assertContains(response, f'<html lang="{language}" dir="{direction}">')
                    text = response.content.decode()
                    self.assertNotIn("%(", text)  # a sign the app should have filled in
                    self.assertNotIn("{{", text)
                    # Digits stay 0 1 2 3 in every language (Arabic-Indic digits would mix with the rest)
                    self.assertIsNone(re.search("[٠-٩۰-۹]", text))

    def test_a_member_sees_her_pages_in_arabic_too(self):
        choose(self.sara, "ar")
        self.login(self.sara)
        for page in ("/", "/my-tasks/", f"/projects/{self.project.pk}/", f"/tasks/{self.task.pk}/",
                     f"/projects/{self.project.pk}/board/", f"/projects/{self.project.pk}/report/", "/dashboard/"):
            with self.subTest(page=page):
                self.assertContains(self.client.get(page), '<html lang="ar" dir="rtl">')

    def test_the_progress_text_of_a_project_is_in_each_language(self):
        project = self.project.pk
        for language, pattern in (("en", r"\d+ of \d+ tasks done"), ("fr", r"\d+ sur \d+ tâches terminées"),
                                  ("ar", r"تم إنجاز \d+ من \d+ (?:مهمة|مهام)")):
            choose(self.amina, language)
            self.login(self.amina)
            for page in (f"/projects/{project}/", f"/projects/{project}/report/"):
                with self.subTest(language=language, page=page):
                    self.assertRegex(self.client.get(page).content.decode(), pattern)

    def test_french_writes_a_space_before_a_colon(self):
        choose(self.sara, "fr")
        self.login(self.sara)
        self.assertContains(self.client.get(f"/activity/?project={self.project.pk}"), "Échéance\xa0:")
        self.assertRegex(self.client.get(f"/projects/{self.project.pk}/report/").content.decode(), r"En cours\xa0: \d")
        choose(self.sara, "ar")  # Arabic and English write the colon right after the word
        self.assertContains(self.client.get(f"/activity/?project={self.project.pk}"), "تاريخ الاستحقاق:")

    def test_dates_are_written_the_way_of_the_language(self):
        choose(self.sara, "fr")
        self.login(self.sara)
        with translation.override("fr"):
            due = date_format(self.task.due_date, gettext("d M Y"))
        self.assertRegex(due, r"^\d{1,2} [a-zé.û]+ \d{4}$")  # like "16 oct. 2026"
        self.assertContains(self.client.get(f"/tasks/{self.task.pk}/"), due)
        # ... and not the English way
        english = date_format(self.task.due_date, "d M Y")
        if english != due:
            self.assertNotContains(self.client.get(f"/tasks/{self.task.pk}/"), english)


class TranslationFilesTests(SimpleTestCase):
    """The French and Arabic files in locale/."""

    def catalog(self, path):
        with open(path, "rb") as file:
            return gettext_module.GNUTranslations(file)._catalog

    def test_the_compiled_files_match_the_po_files(self):
        """The app reads the .mo files. If a .po file is edited and not compiled, people keep seeing
        the old text. Fix: python scripts/compile_translations.py"""
        for code in ("fr", "ar"):
            folder = LOCALE_FOLDER / code / "LC_MESSAGES"
            with tempfile.TemporaryDirectory() as temp:
                fresh = Path(temp) / "django.mo"
                compile_po(folder / "django.po", fresh)
                shipped, expected = self.catalog(folder / "django.mo"), self.catalog(fresh)
            different = sorted((key for key in set(shipped) | set(expected) if shipped.get(key) != expected.get(key)),
                               key=str)
            self.assertEqual(different[:5], [],
                             f"locale/{code}: the .mo is out of date, run scripts/compile_translations.py")

    def test_every_text_is_translated_and_every_sign_is_right(self):
        for code in ("fr", "ar"):
            missing = []
            try:
                read_po(LOCALE_FOLDER / code / "LC_MESSAGES" / "django.po", missing)  # also checks the %(name)s signs
            except TranslationError as error:
                self.fail(str(error))
            self.assertEqual(missing[:5], [], f"locale/{code}: these texts are not translated yet")

    def test_french_and_arabic_have_the_same_texts(self):
        def texts(code):
            catalog = self.catalog(LOCALE_FOLDER / code / "LC_MESSAGES" / "django.mo")
            return {key[0] if isinstance(key, tuple) else key for key in catalog}

        french, arabic = texts("fr"), texts("ar")
        self.assertEqual(sorted(french - arabic)[:5], [], "in French but not in Arabic")
        self.assertEqual(sorted(arabic - french)[:5], [], "in Arabic but not in French")
        self.assertGreater(len(french), 500)

    def test_the_helper_refuses_a_wrong_sign_and_changes_nothing(self):
        with tempfile.TemporaryDirectory() as temp:
            broken = Path(temp) / "django.po"
            broken.write_text(
                'msgid ""\nmsgstr ""\n"Content-Type: text/plain; charset=UTF-8\\n"\n\n'
                '#, python-format\nmsgid "Hello %(name)s,"\nmsgstr "Bonjour %(nom)s,"\n', encoding="utf-8")
            with self.assertRaisesMessage(TranslationError, "nom:s"):
                compile_po(broken, Path(temp) / "django.mo")
            self.assertFalse((Path(temp) / "django.mo").exists())

    def test_arabic_has_six_forms_of_a_counted_word(self):
        # none, one, two, 3 to 10, 11 to 99, and 100 or more
        with translation.override("ar"):
            forms = [ngettext("%(count)d day late", "%(count)d days late", n) for n in (0, 1, 2, 3, 11, 100)]
        self.assertEqual(len(set(forms)), 6)
        self.assertEqual(forms[1:3], ["متأخرة بيوم واحد", "متأخرة بيومين"])

    def test_french_counts_zero_and_one_as_singular(self):
        with translation.override("fr"):
            forms = [ngettext("%(count)d day late", "%(count)d days late", n) for n in (0, 1, 2)]
        self.assertEqual(forms, ["%(count)d jour de retard", "%(count)d jour de retard", "%(count)d jours de retard"])

    def test_the_progress_text_counts_well_in_every_language(self):
        def text(language, done, total):
            with translation.override(language):
                return ngettext("%(done)s of %(counter)s task done", "%(done)s of %(counter)s tasks done",
                                total) % {"done": done, "counter": total}

        self.assertEqual(text("en", 1, 1), "1 of 1 task done")
        self.assertEqual(text("en", 1, 3), "1 of 3 tasks done")
        self.assertEqual([text("fr", 0, 0), text("fr", 1, 1), text("fr", 1, 3)],
                         ["0 sur 0 tâche terminée", "1 sur 1 tâche terminée", "1 sur 3 tâches terminées"])
        # Arabic: none, one, two, 3 to 10, 11 to 99, 100 or more
        self.assertEqual(
            [text("ar", done, total) for done, total in ((0, 0), (1, 1), (1, 2), (1, 3), (5, 11), (5, 100))],
            ["تم إنجاز 0 من 0 مهمة", "تم إنجاز 1 من مهمة واحدة", "تم إنجاز 1 من مهمتين", "تم إنجاز 1 من 3 مهام",
             "تم إنجاز 5 من 11 مهمة", "تم إنجاز 5 من 100 مهمة"])

    def test_the_sentence_about_one_subtask_is_good_french(self):
        def sentence(number):
            return ngettext("Its %(counter)s subtask will be deleted too.",
                            "Its %(counter)s subtasks will be deleted too.", number)

        with translation.override("fr"):
            one, three = sentence(1), sentence(3)
        self.assertEqual(one, "Sa sous-tâche sera aussi supprimée.")  # not "Sa 1 sous-tâche"
        self.assertEqual(three % {"counter": 3}, "Ses 3 sous-tâches seront aussi supprimées.")

    def test_durations_are_written_well_in_each_language(self):
        # "5 minutes", "2 hours, 5 minutes"... come from Django, but our files give the words
        # (Django's own Arabic says "2 دقيقتين", with the number twice)
        now = timezone.now()

        def since(language, **length):
            with translation.override(language):
                # Django keeps the number and its word on the same line with a no-break space: a plain space here
                return timesince(now - timedelta(**length), now).replace("\xa0", " ")

        self.assertEqual(since("en", hours=2, minutes=5), "2 hours, 5 minutes")
        self.assertEqual([since("fr", days=400), since("fr", hours=2, minutes=5)],
                         ["1 an, 1 mois", "2 heures, 5 minutes"])
        self.assertEqual(
            [since("ar", minutes=1), since("ar", minutes=2), since("ar", minutes=5), since("ar", minutes=45),
             since("ar", hours=2, minutes=5), since("ar", days=2), since("ar", days=15), since("ar", days=800)],
            ["دقيقة واحدة", "دقيقتين", "5 دقائق", "45 دقيقة",
             "ساعتين، 5 دقائق", "يومين", "أسبوعين، يوم واحد", "سنتين، شهرين"])
        self.assertEqual(since("ar", seconds=20), "أقل من دقيقة")

    def test_the_words_of_durations_stay_listed_for_translation(self):
        # projects/activity.py lists the six texts Django uses for "5 minutes", so that `makemessages` keeps them
        # in the translation files (where Arabic has better wording). If Django changes one, this says so.
        source = Path(activity.__file__).read_text(encoding="utf-8")
        self.assertEqual(sorted(TIME_STRINGS), ["day", "hour", "minute", "month", "week", "year"])
        for unit in TIME_STRINGS:
            self.assertIn(f'ngettext_lazy("%(num)d {unit}", "%(num)d {unit}s", "num")', source)

    def test_the_feed_says_how_long_ago_in_each_language(self):
        def ago(language, **length):
            event = Event(when=timezone.now() - timedelta(**length), who=None, icon="", color="", action="", name="")
            with translation.override(language):
                return event.ago.replace("\xa0", " ")  # (no-break space, see above)

        self.assertEqual([ago("en", seconds=10), ago("en", minutes=2)], ["just now", "2 minutes ago"])
        self.assertEqual([ago("fr", seconds=10), ago("fr", minutes=2)], ["à l'instant", "il y a 2 minutes"])
        self.assertEqual([ago("ar", seconds=10), ago("ar", minutes=2), ago("ar", hours=3)],
                         ["الآن", "منذ دقيقتين", "منذ 3 ساعات"])

    def test_dates_in_each_language(self):
        july, june = date(2026, 7, 5), date(2026, 6, 5)
        expected = {"en": "05 Jul 2026", "fr": "5 juil. 2026", "ar": "5 يوليو 2026"}
        for code, text in expected.items():
            with translation.override(code):
                self.assertEqual(date_format(july, gettext("d M Y")), text)
        with translation.override("fr"):
            self.assertEqual(date_format(june, gettext("d M Y")), "5 juin 2026")  # not "5 Jui" like July
            self.assertEqual(date_format(july, "l j F Y"), "dimanche 5 juillet 2026")


class EmailLanguageRuleTests(TeamTestCase):
    """Step 23: an email is written in the language of the person's profile. Someone who hasn't chosen a
    language yet gets settings.EMAIL_DEFAULT_LANGUAGE (French)."""

    def assign_vpn_to_sara(self):
        mail.outbox.clear()
        Task.objects.filter(title="Configure the VPN").delete()
        self.client.post(f"/projects/{self.project.pk}/tasks/new/", {
            "title": "Configure the VPN", "status": "todo", "priority": "high",
            "assignee": self.sara.pk, "due_date": days_from_today(4)})
        self.assertEqual(len(mail.outbox), 1)
        return mail.outbox[0]

    def test_the_profile_language_wins(self):
        self.login(self.omar)
        for language, subject in (("ar", "مهمة جديدة لك: Configure the VPN"),
                                  ("en", "New task for you: Configure the VPN"),
                                  ("fr", "Nouvelle tâche pour vous : Configure the VPN")):
            with self.subTest(profile_language=language):
                choose(self.sara, language)
                self.assertEqual(self.assign_vpn_to_sara().subject, subject)

    def test_a_blank_profile_gets_french(self):
        choose(self.sara, "")
        self.login(self.omar)
        email = self.assign_vpn_to_sara()
        self.assertEqual(email.subject, "Nouvelle tâche pour vous : Configure the VPN")
        self.assertIn('<html lang="fr" dir="ltr">', email.alternatives[0][0])

    @override_settings(EMAIL_DEFAULT_LANGUAGE="en")
    def test_the_default_language_can_be_changed(self):
        choose(self.sara, "")
        self.login(self.omar)
        self.assertEqual(self.assign_vpn_to_sara().subject, "New task for you: Configure the VPN")

    @override_settings(EMAIL_DEFAULT_LANGUAGE="klingon")
    def test_an_unknown_default_falls_back_to_english(self):
        choose(self.sara, "")
        self.assertEqual(email_language(User.objects.get(pk=self.sara.pk)), "en")

    def test_the_reminder_follows_the_profile_too(self):
        Task.objects.create(project=self.project, title="Late", assignee=self.sara, due_date=days_from_today(-2))
        for language, start in (("ar", None), ("fr", "Vos tâches"), ("", "Vos tâches"), ("en", "Your tasks")):
            with self.subTest(profile_language=language):
                choose(self.sara, language)
                mail.outbox.clear()
                call_command("send_reminders", "--user", "sara", stdout=io.StringIO())
                email = mail.outbox[0]
                if start:
                    self.assertTrue(email.subject.startswith(start), email.subject)
                else:
                    self.assertIn('dir="rtl"', email.alternatives[0][0])

    def test_the_preview_shows_the_email_as_the_person_will_get_it(self):
        self.login(self.sara)  # logging in saves the language of the login page in the profile...
        choose(self.sara, "")  # ...so the profile is made blank afterwards
        self.assertIn('lang="fr"', self.client.get("/reminders/preview/").content.decode())
