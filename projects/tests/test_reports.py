"""Reports: the Excel files and the printable status report."""

import io

from openpyxl import load_workbook

from projects.models import Project, Task

from .base import TeamTestCase, days_from_today


class ExcelTests(TeamTestCase):
    def open_workbook(self, response):
        self.assertEqual(response.status_code, 200)
        return load_workbook(io.BytesIO(response.content))

    def test_project_tasks_file(self):
        self.login(self.sara)
        response = self.client.get(f"/projects/{self.project.pk}/export.xlsx")
        self.assertIn("attachment;", response["Content-Disposition"])  # a download
        workbook = self.open_workbook(response)
        self.assertEqual(workbook.sheetnames, ["Summary", "Tasks"])

        summary = {row[0]: row[1] for row in workbook["Summary"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual((summary["Project"], summary["Manager"], summary["Tasks"]), ("Firewall", "Omar Ali", 1))

        header, first_task = list(workbook["Tasks"].iter_rows(max_row=2, values_only=True))
        self.assertEqual(header[:3], ("Task", "Main task", "Status"))
        self.assertEqual(first_task[0], "Install the firewall")
        self.assertEqual(first_task[6].date(), days_from_today(5))  # a real date that Excel can sort

    def test_text_that_looks_like_a_formula_stays_text(self):
        # Excel runs anything that starts with "=" as a formula: a title must never do that
        Task.objects.create(project=self.project, title='=HYPERLINK("http://evil.example","Click")')
        self.login(self.sara)
        sheet = self.open_workbook(self.client.get(f"/projects/{self.project.pk}/export.xlsx"))["Tasks"]
        cell = sheet["A3"]
        self.assertEqual(cell.value, '=HYPERLINK("http://evil.example","Click")')
        self.assertEqual(cell.data_type, "s")  # "s" = string (text), not "f" = formula

    def test_project_list_file_uses_the_page_filters(self):
        Project.objects.create(name="Old laptops", manager=self.omar, status="completed")
        self.login(self.sara)
        workbook = self.open_workbook(self.client.get("/projects/export.xlsx?status=completed"))
        names = [row[0] for row in workbook["Projects"].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(names, ["Old laptops"])


class StatusReportTests(TeamTestCase):
    def test_report_page(self):
        Task.objects.create(project=self.project, title="Order", status="done")
        Task.objects.create(project=self.project, title="Late one", due_date=days_from_today(-3))
        self.login(self.sara)
        response = self.client.get(f"/projects/{self.project.pk}/report/")
        self.assertContains(response, "33% · 1 of 3 tasks done")
        self.assertEqual([task.title for task in response.context["overdue"]], ["Late one"])
        self.assertContains(response, "3 days")  # late by
        self.assertContains(response, "window.print()")  # the Print button
