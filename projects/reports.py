"""Reports: Excel files to download, and a status report page made for printing.
Everyone in a project's team can open its reports (they only read). Admins can open all of them."""
from datetime import date

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, render
from django.utils import timezone
from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Task
from .views import filter_projects, visible_projects

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="206BC4")  # the app's blue
DATE_FORMAT = "dd mmm yyyy"


def add_table_sheet(workbook, title, headers, rows, widths):
    """A sheet holding a table: a colored header row that stays visible when scrolling,
    a filter button on every column, dates shown as dates, and text always kept as text."""
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for row in rows:
        sheet.append(row)
    for row_cells in sheet.iter_rows(min_row=2):
        for cell in row_cells:
            if isinstance(cell.value, date):
                cell.number_format = DATE_FORMAT
            elif isinstance(cell.value, str) and cell.value.startswith("="):
                # A title like "=HYPERLINK(...)" must stay text: Excel would run it as a formula
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    return sheet


def link_first_column(sheet, request, objects):
    """Make each name in column A a link to its page in the app."""
    for row, obj in enumerate(objects, start=2):
        cell = sheet.cell(row=row, column=1)
        cell.hyperlink = request.build_absolute_uri(obj.get_absolute_url())
        cell.style = "Hyperlink"


def excel_response(workbook, filename):
    """Send the workbook to the browser as a file to save."""
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    workbook.save(response)
    return response


def ordered_tasks(project):
    """The project's tasks in reading order: each main task, then its subtasks. Sorted by due date."""
    tasks = list(
        project.tasks.select_related("assignee", "parent")
        .prefetch_related("depends_on")
        .order_by("due_date", "id")
    )
    ordered = []
    for main in (t for t in tasks if t.parent_id is None):
        ordered.append(main)
        ordered.extend(t for t in tasks if t.parent_id == main.pk)
    return ordered


def done_on(task):
    """The day a task was finished, in the local time zone (Excel can't store time zones)."""
    return timezone.localtime(task.completed_at).date() if task.completed_at else None


@login_required
def project_tasks_excel(request, pk):
    """An Excel file with a summary of the project and all its tasks."""
    project = get_object_or_404(visible_projects(request).with_task_counts().select_related("manager"), pk=pk)
    tasks = ordered_tasks(project)
    today = timezone.localdate()

    workbook = Workbook()
    workbook.remove(workbook.active)  # start without the empty default sheet

    summary = add_table_sheet(workbook, "Summary", ["Item", "Value"], [
        ["Project", project.name],
        ["Status", project.get_status_display()],
        ["Priority", project.get_priority_display()],
        ["Manager", str(project.manager)],
        ["Start date", project.start_date],
        ["End date", project.end_date],
        ["Progress (%)", project.progress],
        ["Tasks", project.task_total],
        ["Tasks done", project.task_done],
        ["Overdue tasks", project.task_overdue],
        ["Exported on", today],
        ["Exported by", str(request.user)],
    ], widths=[18, 45])
    summary.auto_filter.ref = None  # no filter buttons on a two-column summary

    sheet = add_table_sheet(
        workbook, "Tasks",
        ["Task", "Main task", "Status", "Priority", "Assigned to", "Start date", "Due date",
         "Done on", "Milestone", "Waits for", "Late"],
        [[
            task.title,
            task.parent.title if task.parent else "",
            task.get_status_display(),
            task.get_priority_display(),
            str(task.assignee) if task.assignee else "Unassigned",
            task.start_date,
            task.due_date,
            done_on(task),
            "Yes" if task.is_milestone else "",
            ", ".join(other.title for other in task.depends_on.all()),
            "Yes" if task.is_overdue else "",
        ] for task in tasks],
        widths=[40, 28, 13, 11, 20, 13, 13, 13, 10, 35, 7],
    )
    link_first_column(sheet, request, tasks)

    filename = f"{slugify(project.name) or 'project'}-tasks-{today:%Y-%m-%d}.xlsx"
    return excel_response(workbook, filename)


@login_required
def projects_excel(request):
    """The project list as an Excel file, with the same search and filter as the page."""
    projects, status, query = filter_projects(request)
    projects = list(projects)
    today = timezone.localdate()

    workbook = Workbook()
    workbook.remove(workbook.active)
    sheet = add_table_sheet(
        workbook, "Projects",
        ["Project", "Status", "Priority", "Manager", "Start date", "End date",
         "Tasks", "Done", "Overdue tasks", "Progress (%)", "Late"],
        [[
            project.name,
            project.get_status_display(),
            project.get_priority_display(),
            str(project.manager),
            project.start_date,
            project.end_date,
            project.task_total,
            project.task_done,
            project.task_overdue,
            project.progress,
            "Yes" if project.is_overdue else "",
        ] for project in projects],
        widths=[40, 13, 11, 20, 13, 13, 8, 8, 14, 13, 7],
    )
    link_first_column(sheet, request, projects)
    return excel_response(workbook, f"projects-{today:%Y-%m-%d}.xlsx")


@login_required
def project_report(request, pk):
    """A one-page status report of a project, laid out for printing or saving as a PDF."""
    project = get_object_or_404(visible_projects(request).with_task_counts().select_related("manager"), pk=pk)
    tasks = ordered_tasks(project)
    today = timezone.localdate()

    counts = {status: 0 for status in Task.Status.values}
    for task in tasks:
        counts[task.status] += 1

    overdue = sorted((t for t in tasks if t.is_overdue), key=lambda t: t.due_date)
    for task in overdue:
        task.days_late = (today - task.due_date).days

    context = {
        "project": project,
        "tasks": tasks,
        "status_counts": [
            (label, counts[status], Task.STATUS_COLORS[status]) for status, label in Task.Status.choices
        ],
        "milestones": [t for t in tasks if t.is_milestone],
        "overdue": overdue,
        "generated_at": timezone.localtime(),
    }
    return render(request, "projects/project_report.html", context)
